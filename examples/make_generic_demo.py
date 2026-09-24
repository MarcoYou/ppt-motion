#!/usr/bin/env python3
"""Create an independent four-slide fixture for PPT Motion.

Usage:
    python3 examples/make_generic_demo.py --out /tmp/generic-source

Requires PyMuPDF and python-pptx (including its xlsxwriter dependency):
    python3 -m pip install pymupdf python-pptx

The PDF is deliberately drawn with explicit geometry. The PPTX contains native
charts/tables with the same data and matching object bounds, but it is NOT an
export of that PPTX. Font metrics and Office chart layout can differ. The PDF
is the geometry reference for the motion fixture; inspect an actual PowerPoint
export before claiming renderer-level equivalence.

All measurements in demo-layout.json use PDF points, top-left origin, and
[x0, y0, x1, y1] boxes on a 960 x 540 canvas. The data is entirely synthetic.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

try:
    import fitz
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_MARK
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Pt
except ImportError as exc:
    raise SystemExit(
        f"Missing fixture dependency: {exc.name}. "
        "Use a Python environment with pymupdf and python-pptx installed.\n"
        "Install there with: python3 -m pip install pymupdf python-pptx"
    ) from exc

WIDTH, HEIGHT = 960, 540
INK = "203444"
MUTED = "596A74"
BLUE = "3476A1"
RED = "B86159"
GRID = "DCE3E7"
WHITE = "FFFFFF"
PALE = "F3F6F7"
FONT = "Arial"


def rgb(color: str) -> tuple[float, float, float]:
    return tuple(int(color[i:i + 2], 16) / 255 for i in (0, 2, 4))


def rect(box: list[float]) -> fitz.Rect:
    return fitz.Rect(*box)


def ppt_box(box: list[float]) -> tuple[Any, Any, Any, Any]:
    x0, y0, x1, y1 = box
    return Pt(x0), Pt(y0), Pt(x1 - x0), Pt(y1 - y0)


def text(
    pdf: Any, slide: Any, name: str, content: str, box: list[float],
    size: float = 18, color: str = INK, bold: bool = False,
    align: str = "left", valign: str = "top",
) -> None:
    """Write the same text and outer bounds into both sources."""
    pdf_font = "hebo" if bold else "helv"
    pdf_align = {"left": 0, "center": 1, "right": 2}[align]
    text_box = list(box)
    if valign == "middle":
        lines = content.count("\n") + 1
        text_box[1] += max(0, ((box[3] - box[1]) - lines * size * 1.2) / 2)
    spare = pdf.insert_textbox(
        rect(text_box), content, fontname=pdf_font, fontsize=size,
        color=rgb(color), align=pdf_align, lineheight=1.18,
    )
    if spare < -0.1:
        raise ValueError(f"PDF text overflow in {name}: {spare:.2f} points")
    shape = slide.shapes.add_textbox(*ppt_box(box))
    shape.name = name
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if valign == "middle" else MSO_ANCHOR.TOP
    for index, line in enumerate(content.split("\n")):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.text = line
        paragraph.alignment = {
            "left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT,
        }[align]
        paragraph.space_after = Pt(0)
        paragraph.space_before = Pt(0)
        paragraph.line_spacing = 1.18
        for run in paragraph.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor.from_string(color)


def pdf_text(
    pdf: Any, content: str, box: list[float], size: float = 13,
    color: str = MUTED, align: int = 0, bold: bool = False,
) -> None:
    spare = pdf.insert_textbox(
        rect(box), content, fontname="hebo" if bold else "helv", fontsize=size,
        color=rgb(color), align=align, lineheight=1.1,
    )
    if spare < -0.1:
        raise ValueError(f"PDF text overflow: {content!r}")


def make_page(pdf_doc: Any, prs: Any, page_number: int, title: str) -> tuple[Any, Any]:
    pdf = pdf_doc.new_page(width=WIDTH, height=HEIGHT)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    text(pdf, slide, "slide.title", title, [48, 28, 912, 77], 30, bold=True)
    text(
        pdf, slide, "slide.source-note", "Illustrative data for a reusable motion example",
        [48, 503, 845, 523], 11, MUTED,
    )
    text(pdf, slide, "slide.page-number", str(page_number), [875, 503, 912, 523], 11, MUTED, align="right")
    return pdf, slide


def add_region(page: dict[str, Any], name: str, role: str, box: list[float], **metadata: Any) -> None:
    page["namedregions"][name] = {"role": role, "box": box, **metadata}


def page_metadata(page_number: int, title: str, message: str) -> dict[str, Any]:
    result: dict[str, Any] = {"page": page_number, "title": title, "message": message, "namedregions": {}}
    add_region(result, "title", "title", [48, 28, 912, 77], pptx_object="slide.title", recommended="static")
    return result


def chart_layout(chart: Any, outer: list[float], plot: list[float]) -> None:
    """Request an inner plot area near the explicit PDF plot; Office may adjust."""
    from pptx.oxml.ns import qn

    plot_area = chart._chartSpace.find(".//" + qn("c:plotArea"))
    layout = plot_area.find(qn("c:layout"))
    if layout is None:
        layout = OxmlElement("c:layout")
        plot_area.insert(0, layout)
    manual = OxmlElement("c:manualLayout")
    values = [
        ("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"),
        ("wMode", "factor"), ("hMode", "factor"),
        ("x", (plot[0] - outer[0]) / (outer[2] - outer[0])),
        ("y", (plot[1] - outer[1]) / (outer[3] - outer[1])),
        ("w", (plot[2] - plot[0]) / (outer[2] - outer[0])),
        ("h", (plot[3] - plot[1]) / (outer[3] - outer[1])),
    ]
    for tag, value in values:
        element = OxmlElement("c:" + tag)
        element.set("val", str(value))
        manual.append(element)
    layout.append(manual)


def format_chart(chart: Any, minimum: float, maximum: float, major: float) -> None:
    chart.has_title = False
    chart.has_legend = False
    chart.font.name = FONT
    chart.font.size = Pt(13)
    chart.font.color.rgb = RGBColor.from_string(MUTED)
    value_axis = chart.value_axis
    value_axis.minimum_scale = minimum
    value_axis.maximum_scale = maximum
    value_axis.major_unit = major
    value_axis.has_major_gridlines = True
    value_axis.major_gridlines.format.line.color.rgb = RGBColor.from_string(GRID)
    value_axis.major_gridlines.format.line.width = Pt(0.6)
    value_axis.tick_labels.number_format = "0"
    for axis in (value_axis, chart.category_axis):
        axis.major_tick_mark = XL_TICK_MARK.NONE
        axis.minor_tick_mark = XL_TICK_MARK.NONE
        axis.tick_labels.font.name = FONT
        axis.tick_labels.font.size = Pt(13)
        axis.tick_labels.font.color.rgb = RGBColor.from_string(MUTED)
        axis.format.line.color.rgb = RGBColor.from_string(GRID)
        axis.format.line.width = Pt(0.6)


def pdf_axes(
    pdf: Any, plot: list[float], minimum: float, maximum: float,
    ticks: list[float], categories: list[str], centers: list[float],
) -> None:
    x0, y0, x1, y1 = plot
    for value in ticks:
        y = y1 - (value - minimum) / (maximum - minimum) * (y1 - y0)
        pdf.draw_line((x0, y), (x1, y), color=rgb(GRID), width=0.7)
        pdf_text(pdf, f"{value:g}", [x0 - 45, y - 8, x0 - 12, y + 12], 12, align=2)
    for label, x in zip(categories, centers):
        pdf_text(pdf, label, [x - 55, y1 + 13, x + 55, y1 + 36], 13, align=1)


def native_table(
    slide: Any, name: str, values: list[list[str]], box: list[float],
    fills: list[list[str]], colors: list[list[str]], size: float = 16,
    header: bool = False,
) -> Any:
    rows, cols = len(values), len(values[0])
    shape = slide.shapes.add_table(rows, cols, *ppt_box(box))
    shape.name = name
    table = shape.table
    table.first_row = header
    table.horz_banding = False
    table.vert_banding = False
    for row_index, row in enumerate(values):
        table.rows[row_index].height = Pt((box[3] - box[1]) / rows)
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = value
            cell.margin_left = cell.margin_right = Pt(9)
            cell.margin_top = cell.margin_bottom = Pt(2)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(fills[row_index][col_index])
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(size)
                    run.font.bold = header and row_index == 0
                    run.font.color.rgb = RGBColor.from_string(colors[row_index][col_index])
    return shape


def cover(pdf_doc: Any, prs: Any) -> dict[str, Any]:
    title = "Research motion demo"
    message = "A small independent deck exercises text, bars, lines and heatmap tiles."
    pdf, slide = make_page(pdf_doc, prs, 1, title)
    text(pdf, slide, "cover.headline", "A reusable presentation example", [48, 167, 900, 225], 35, bold=True)
    body = "Four slides with editable PowerPoint content.\nThe PDF supplies explicit geometry for motion selection.\nAll values are synthetic."
    text(pdf, slide, "cover.body", body, [48, 252, 860, 365], 21, MUTED)
    metadata = page_metadata(1, title, message)
    add_region(metadata, "headline", "text", [48, 167, 900, 225], pptx_object="cover.headline", recommended="static")
    add_region(metadata, "body", "text", [48, 252, 860, 365], pptx_object="cover.body", recommended="static")
    return metadata


def bars_and_table(pdf_doc: Any, prs: Any) -> dict[str, Any]:
    title = "Contributions can be positive or negative"
    message = "Two products add to the total while Product B subtracts from it."
    pdf, slide = make_page(pdf_doc, prs, 2, title)
    text(pdf, slide, "bars.subtitle", "Contribution to the quarterly total", [48, 89, 880, 122], 18, MUTED)
    outer = [60, 142, 565, 444]
    plot = [104, 170, 536, 410]
    zero = 330
    minimum, maximum = -4, 8
    categories = ["Product A", "Product B", "Product C"]
    values = [6, -3, 4.5]
    centers = [176, 320, 464]
    pdf_axes(pdf, plot, minimum, maximum, [-4, 0, 4, 8], categories, centers)
    pdf.draw_line((plot[0], zero), (plot[2], zero), color=rgb(MUTED), width=1)
    items = []
    for label, value, center in zip(categories, values, centers):
        end = zero - value * 20
        box = [center - 35, min(zero, end), center + 35, max(zero, end)]
        color = BLUE if value >= 0 else RED
        pdf.draw_rect(rect(box), color=None, fill=rgb(color), width=0)
        label_box = [center - 35, end - 27, center + 35, end - 7] if value > 0 else [center - 35, end + 6, center + 35, end + 26]
        pdf_text(pdf, f"{value:+g}", label_box, 13, INK, align=1)
        items.append({"name": label, "value": value, "box": box, "origin": [center, zero], "fill": "#" + color.lower()})
    chart_data = CategoryChartData()
    chart_data.categories = categories
    chart_data.add_series("Contribution", values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, *ppt_box(outer), chart_data)
    frame.name = "chart.contribution-bars"
    chart = frame.chart
    format_chart(chart, minimum, maximum, 4)
    chart_layout(chart, outer, plot)
    chart.plots[0].gap_width = 105
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    chart.plots[0].data_labels.number_format = "+0.#;-0.#;0"
    chart.plots[0].data_labels.font.size = Pt(13)
    for point, value in zip(chart.series[0].points, values):
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RGBColor.from_string(BLUE if value >= 0 else RED)
        point.format.line.fill.background()
    table_box = [610, 170, 904, 346]
    table_values = [["Product", "Contribution"], ["A", "+6.0"], ["B", "-3.0"], ["C", "+4.5"]]
    fills = [[INK, INK]] + [[PALE, PALE], [WHITE, WHITE], [PALE, PALE]]
    colors = [[WHITE, WHITE]] + [[INK, INK] for _ in range(3)]
    native_table(slide, "table.contributions", table_values, table_box, fills, colors, 16, header=True)
    cells = []
    for row in range(4):
        for col in range(2):
            box = [610 + col * 147, 170 + row * 44, 610 + (col + 1) * 147, 170 + (row + 1) * 44]
            pdf.draw_rect(rect(box), color=rgb(WHITE), fill=rgb(fills[row][col]), width=1)
            pdf_text(pdf, table_values[row][col], [box[0] + 8, box[1] + 12, box[2] - 8, box[3] - 5], 16, colors[row][col], align=1, bold=row == 0)
            cells.append({"row": row, "column": col, "text": table_values[row][col], "box": box})
    text(pdf, slide, "bars.takeaway", "Both directions start at the same zero.", [610, 375, 904, 440], 18)
    metadata = page_metadata(2, title, message)
    add_region(metadata, "bars", "bars", plot, pptx_object=frame.name, orientation="vertical", axis="y", zero=zero, range=[minimum, maximum], values=values, categories=categories, items=items, recommended="bar-y")
    add_region(metadata, "table", "table", table_box, pptx_object="table.contributions", rows=4, columns=2, cells=cells, recommended="static")
    add_region(metadata, "takeaway", "text", [610, 375, 904, 440], pptx_object="bars.takeaway", recommended="static")
    return metadata


def line_chart(pdf_doc: Any, prs: Any) -> dict[str, Any]:
    title = "The series finishes above its starting level"
    message = "The final value is higher despite a mid-period decline."
    pdf, slide = make_page(pdf_doc, prs, 3, title)
    text(pdf, slide, "line.subtitle", "Illustrative index level", [48, 88, 880, 121], 18, MUTED)
    outer = [60, 132, 920, 430]
    plot = [105, 160, 884, 390]
    categories = ["Q1", "Q2", "Q3", "Q4", "Q5", "Q6"]
    values = [18, 24, 21, 32, 38, 45]
    centers = [plot[0] + i * (plot[2] - plot[0]) / 5 for i in range(6)]
    points = [[x, plot[3] - value / 50 * (plot[3] - plot[1])] for x, value in zip(centers, values)]
    pdf_axes(pdf, plot, 0, 50, [0, 10, 20, 30, 40, 50], categories, centers)
    pdf.draw_polyline(points, color=rgb(BLUE), width=3)
    data = CategoryChartData()
    data.categories = categories
    data.add_series("Index", values)
    frame = slide.shapes.add_chart(XL_CHART_TYPE.LINE, *ppt_box(outer), data)
    frame.name = "chart.index-line"
    chart = frame.chart
    format_chart(chart, 0, 50, 10)
    chart_layout(chart, outer, plot)
    chart.series[0].format.line.color.rgb = RGBColor.from_string(BLUE)
    chart.series[0].format.line.width = Pt(3)
    chart.series[0].smooth = False
    body = "The level rises from 18 to 45, with a brief decline in Q3."
    text(pdf, slide, "line.takeaway", body, [70, 450, 900, 483], 20)
    metadata = page_metadata(3, title, message)
    add_region(metadata, "plot", "chart", plot, pptx_object=frame.name, range=[0, 50], categories=categories, values=values, recommended="static")
    line_box = [min(p[0] for p in points), min(p[1] for p in points), max(p[0] for p in points), max(p[1] for p in points)]
    add_region(metadata, "line", "line", line_box, pptx_object=frame.name, points=points, stroke="#" + BLUE.lower(), stroke_width=3, recommended="line")
    add_region(metadata, "takeaway", "text", [70, 450, 900, 483], pptx_object="line.takeaway", recommended="static")
    return metadata


def heat_color(value: float) -> str:
    neutral = (237, 241, 236)
    end = (62, 126, 91) if value >= 0 else (185, 91, 79)
    fraction = value / 8 if value >= 0 else -value / 4
    return "".join(f"{round(a + fraction * (b - a)):02X}" for a, b in zip(neutral, end))


def heatmap(pdf_doc: Any, prs: Any) -> dict[str, Any]:
    title = "Improvement broadens over time"
    message = "All four segments show positive values in the last three years."
    pdf, slide = make_page(pdf_doc, prs, 4, title)
    text(pdf, slide, "heatmap.subtitle", "Illustrative segment scores", [48, 88, 880, 121], 18, MUTED)
    box = [180, 170, 780, 390]
    values = [[-3, -1, 2, 5, 7], [-4, -2, 1, 4, 6], [-1, 0, 3, 6, 8], [-2, -1, 2, 4, 5]]
    row_names = ["Segment A", "Segment B", "Segment C", "Segment D"]
    col_names = ["Year 1", "Year 2", "Year 3", "Year 4", "Year 5"]
    fills = [[heat_color(value) for value in row] for row in values]
    colors = [[WHITE if abs(value) >= (3 if value < 0 else 6) else INK for value in row] for row in values]
    native_table(slide, "table.heatmap", [[f"{v:+g}" if v != 0 else "0" for v in row] for row in values], box, fills, colors, 19)
    cells = []
    for row, label in enumerate(row_names):
        text(pdf, slide, f"heatmap.row-{row + 1}", label, [48, 170 + row * 55, 162, 225 + row * 55], 16, MUTED, align="right", valign="middle")
    for col, label in enumerate(col_names):
        text(pdf, slide, f"heatmap.column-{col + 1}", label, [180 + col * 120, 134, 300 + col * 120, 161], 15, MUTED, align="center")
    for row in range(4):
        for col in range(5):
            cell_box = [180 + col * 120, 170 + row * 55, 300 + col * 120, 225 + row * 55]
            pdf.draw_rect(rect(cell_box), color=rgb(WHITE), fill=rgb(fills[row][col]), width=1.5)
            value = values[row][col]
            label = f"{value:+g}" if value else "0"
            pdf_text(pdf, label, [cell_box[0] + 5, cell_box[1] + 17, cell_box[2] - 5, cell_box[3] - 7], 19, colors[row][col], align=1)
            cells.append({"row": row, "column": col, "value": value, "box": cell_box, "fill": "#" + fills[row][col].lower()})
    text(pdf, slide, "heatmap.takeaway", "Every segment has a positive score by Year 3.", [70, 447, 900, 480], 20)
    metadata = page_metadata(4, title, message)
    add_region(metadata, "heatmap", "heatmap", box, pptx_object="table.heatmap", rows=4, columns=5, row_labels=row_names, column_labels=col_names, values=values, cells=cells, recommended="tile")
    add_region(metadata, "positive-years", "emphasis", [420, 170, 780, 390], recommended="outline", source_annotation=False)
    add_region(metadata, "takeaway", "text", [70, 447, 900, 480], pptx_object="heatmap.takeaway", recommended="static")
    return metadata


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", type=Path, required=True, help="Directory for generic-demo.pdf, generic-demo.pptx and demo-layout.json")
    parser.add_argument("--force", action="store_true", help="Replace these generated fixture files if they already exist")
    args = parser.parse_args()
    output = args.out.expanduser().resolve()
    paths = {"pdf": output / "generic-demo.pdf", "pptx": output / "generic-demo.pptx", "layout": output / "demo-layout.json"}
    existing = [str(path) for path in paths.values() if path.exists()]
    if existing and not args.force:
        parser.error("Generated file(s) already exist. Use a new --out or --force:\n" + "\n".join(existing))
    output.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(WIDTH), Pt(HEIGHT)
    prs.core_properties.title = "Research motion demo"
    prs.core_properties.subject = "Independent synthetic PPT Motion fixture"
    prs.core_properties.author = "PPT Motion example"
    prs.core_properties.keywords = "synthetic, bars, line, heatmap, native charts"
    doc = fitz.open()
    pages = [cover(doc, prs), bars_and_table(doc, prs), line_chart(doc, prs), heatmap(doc, prs)]
    doc.set_metadata({"title": "Research motion demo", "author": "PPT Motion example", "subject": "Synthetic data with explicit fixture geometry"})
    doc.save(str(paths["pdf"]), garbage=4, deflate=True)
    doc.close()
    prs.save(str(paths["pptx"]))
    layout = {
        "schema_version": 1,
        "fixture": "generic-research-demo",
        "data_status": "synthetic",
        "canvas": {"width": WIDTH, "height": HEIGHT, "units": "pt", "origin": "top-left", "box_format": "x0,y0,x1,y1"},
        "sources": {"pdf": paths["pdf"].name, "pptx": paths["pptx"].name},
        "rendering_note": "The PDF is drawn independently. PPTX object bounds/content match the fixture design, but native chart layout and font metrics can vary by renderer. This is not a certified matching PowerPoint PDF export.",
        "pages": pages,
    }
    paths["layout"].write_text(json.dumps(layout, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"pages": len(pages), "files": {key: str(value) for key, value in paths.items()}, "bar_zero_y": 330, "heatmap_cells": 20}, indent=2))


if __name__ == "__main__":
    main()
